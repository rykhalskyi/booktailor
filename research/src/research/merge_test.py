from docxcompose.composer import Composer
from docx import Document

def merge_documents(doc1_path, doc2_path, output_path):
    master = Document(doc1_path)
    composer = Composer(master)
    doc2 = Document(doc2_path)
    composer.append(doc2)
    composer.save(output_path)


#master = Document("../../docx/PartA.docx")
#composer = Composer(master)
#doc1 = Document("../../docx/PartB.docx")
#composer.append(doc1)
#composer.save("../../docx/combined.docx")